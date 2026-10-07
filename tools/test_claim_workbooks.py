#!/usr/bin/env python3
"""test_claim_workbooks.py -- claim_workbooks.py appends to the two shared workbooks, never rebuilds them.

On small synthetic workbooks shaped like the real ones (the builders' sheets, headers and styles):
  1. claims-append: one new row per held MRI session not yet listed, keyed as the builder keyed it
     (the study folder; a `<study>__<exam>` name joins its study); new acquisitions of a listed
     session go to "Acquisitions" only; an Added column (date + source) and one Read me line; the
     autofilter and the validation extended; every existing cell unchanged; a re-run appends nothing.
  2. assign-append: the drive-3 holding files and blank-project acquisitions grouped by the 2b
     functions themselves, so the group keys written are the ones `remap` / `apply-raw` join on
     (proved by running nonraw_placement.remap_rows on the answer); numbering after the last G.
  3. Refusals and safety: Excel's lock file, an existing or on-NAS --backup-dir, --apply without
     one, an edited header, a held non-MRI row; the dry run writes only a preview off the NAS;
     the verification catches a changed cell, and a failed final check puts the copy back.

Temporary directories only: no NAS, no network.
Run:  python tools/test_claim_workbooks.py      (exit 0 = pass)
"""

import contextlib
import csv
import datetime as dt
import hashlib
import io
import os
import sys
import tempfile

_THIS_DIR = os.path.dirname(os.path.abspath(__file__))
for _p in (_THIS_DIR, os.path.join(_THIS_DIR, "drive_staging")):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

import claim_workbooks as cw  # noqa: E402
import nonraw_placement as NP  # noqa: E402
from ingest import projects_registry, registry  # noqa: E402

HOLD = registry.OPERATOR_HOLD
TODAY = "2026-10-07"
FAILS = []


def check(cond, msg):
    if not cond:
        FAILS.append(msg)
        print(f"  FAIL: {msg}")
    else:
        print(f"  ok:   {msg}")


def run(*args):
    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        try:
            rc = cw.main(list(args))
        except SystemExit as exc:
            rc = exc.code
    return rc, out.getvalue(), err.getvalue()


def sha(path):
    with open(path, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


@contextlib.contextmanager
def patched(obj, name, value):
    old = getattr(obj, name)
    setattr(obj, name, value)
    try:
        yield
    finally:
        setattr(obj, name, old)


# ---- Fixtures ----------------------------------------------------------------------

def reg_row(acq, **kw):
    row = dict.fromkeys(registry.REGISTRY_FIELDS, "")
    row.update({"acq_id": acq, "instrument": "MRI", "data_source": "internal",
                "acquisition_datetime": f"{acq[4:8]}-{acq[8:10]}-{acq[10:12]}T10:21:42.100+01:00"})
    row.update(kw)
    return row


def write_nas(d, rows, projects=()):
    nas = os.path.join(d, "nas")
    os.makedirs(os.path.join(nas, "registries"))
    os.makedirs(os.path.join(nas, "projects"))
    with open(os.path.join(nas, "registries", "registry_raw.csv"), "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\r\n")
        w.writerow(registry.REGISTRY_FIELDS)
        for r in rows:
            w.writerow([r[k] for k in registry.REGISTRY_FIELDS])
    with open(os.path.join(nas, "registries", "registry_projects.csv"), "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=projects_registry.PROJECT_REGISTRY_FIELDS, extrasaction="ignore")
        w.writeheader()
        for p in projects:
            w.writerow(p)
    return nas


HEAD = PatternFill("solid", fgColor="1F4E78")
FILL = PatternFill("solid", fgColor="FFF2CC")
S1 = "20220124_083002_jrc240122_m23_0219_1_1"


def claims_workbook(path):
    """The builder's four sheets, one listed session S1 (acquisitions A1, A2)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    for i, t in enumerate(["MRI sessions: who ran the scanner?", "", "Questions: the Data Office (Ryan Tasseff)."], 1):
        ws.cell(row=i, column=1, value=t or None)
    ws = wb.create_sheet("Sessions to claim")
    cols = ["Project", "Date", "Start", "Session folder on the scanner", "Animal", "Acquisitions", "Who ran it",
            "Notes", "Answered by", "ACQ-IDs (first … last)", "session key (do not edit)"]
    for j, n in enumerate(cols, 1):
        c = ws.cell(row=1, column=j, value=n)
        c.fill, c.font = HEAD, Font(bold=True, color="FFFFFF")
    vals = ["AE-biomaGUNE-0219", dt.datetime(2022, 1, 24), "08:30", S1, "m23_0219", 2, "Irene", None, "Irene",
            "ACQ-20220124-MRI-001 … ACQ-20220124-MRI-002", S1]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=2, column=j, value=v)
        if j == 2:
            c.number_format = "yyyy-mm-dd"
        if j in (7, 8, 9):
            c.fill = FILL
    ws.auto_filter.ref = "A1:K2"
    dv = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="200", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("G2:G2")
    wp = wb.create_sheet("Projects")
    for j, n in enumerate(["Project", "Status", "Sessions", "Acquisitions", "First session", "Last session"], 1):
        wp.cell(row=1, column=j, value=n)
    for j, v in enumerate(["AE-biomaGUNE-0219", "active", 1, 2], 1):
        wp.cell(row=2, column=j, value=v)
    wp.cell(row=3, column=1, value="Total")
    wa = wb.create_sheet("Acquisitions")
    for j, n in enumerate(["ACQ-ID", "Session folder on the scanner", "Exam", "Acquired", "Project", "Animal",
                           "Subject (facility id)"], 1):
        wa.cell(row=1, column=j, value=n)
    for i, (acq, exam) in enumerate((("ACQ-20220124-MRI-001", "1"), ("ACQ-20220124-MRI-002", "2")), 2):
        for j, v in enumerate([acq, S1, exam, "2022-01-24 08:45:07", "AE-biomaGUNE-0219", "m23_0219",
                               "23-AE-biomaGUNE-0219"], 1):
            wa.cell(row=i, column=j, value=v)
    wa.auto_filter.ref = "A1:G3"
    wb.save(path)


def claims_rows():
    return [
        reg_row("ACQ-20220124-MRI-001", operator=HOLD, original_name=S1 + "/1", project_id="PROJ-0010"),
        reg_row("ACQ-20220124-MRI-002", operator=HOLD, original_name=S1 + "/2", project_id="PROJ-0010"),
        reg_row("ACQ-20220124-MRI-003", operator=HOLD, original_name=S1 + "/7", project_id="PROJ-0010",
                ingest_config="tools/configs/drives_2026-09/dicom/b.yaml"),            # new acq, listed session
        reg_row("ACQ-20220124-MRI-004", operator=HOLD, original_name=S1 + "__9", project_id="PROJ-0010",
                ingest_config="tools/configs/drives_2026-09/dicom/x1.yaml"),           # no '/': joins S1
        reg_row("ACQ-20190909-MRI-001", operator=HOLD, original_name="20190909_124235_jrc190909_r111_2DG2_1_1/1",
                project_id="PROJ-0060", sample_id="r111_0118",
                ingest_config="tools/configs/drive3_mri/drive3_mri_M01_0118.yaml"),
        reg_row("ACQ-20190909-MRI-002", operator=HOLD, original_name="20190909_124235_jrc190909_r111_2DG2_1_1/20",
                project_id="PROJ-0060", sample_id="r111_0118",
                ingest_config="tools\\configs\\drive3_mri\\drive3_mri_M01_0118.yaml"),
        reg_row("ACQ-20210420-MRI-001", operator=HOLD, original_name="20210420_155137_jrc210420_m62_1019_1_1/1",
                sample_id="m62_1019", ingest_config="tools/configs/mri_1019_kgjesus_2021_mes02.yaml"),
        reg_row("ACQ-20200101-MRI-001", operator="", original_name="20200101_100000_jrc_x_1_1/1"),   # blank: not held
        reg_row("ACQ-20200101-MRI-002", operator="Irene", original_name="20200101_100000_jrc_y_1_1/1"),
    ]


PROJECTS = [{"project_id": "PROJ-0010", "name": "AE-biomaGUNE-0219", "status": "active"},
            {"project_id": "PROJ-0060", "name": "AE-biomaGUNE-0118", "status": "active"},
            {"project_id": "PROJ-0004", "name": "AE-biomaGUNE-0619", "status": "active",
             "folder_location": "/projects/AE-biomaGUNE-0619/"}]


def cells(path):
    """{(sheet, row, col): (value, number_format)} of every non-empty cell."""
    wb = load_workbook(path)
    out = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    out[(ws.title, c.row, c.column)] = (c.value, c.number_format)
    return out


def setup_claims(d, rows=None):
    nas = write_nas(d, claims_rows() if rows is None else rows, PROJECTS)
    wbp = os.path.join(nas, "projects", cw.CLAIMS_NAME)
    claims_workbook(wbp)
    return nas, wbp


# ---- 1. claims-append ----------------------------------------------------------------

def test_claims_dry_run():
    print("claims-append: the dry run writes only a preview, off the NAS:")
    with tempfile.TemporaryDirectory() as d:
        nas, wbp = setup_claims(d)
        before, mtime = sha(wbp), os.stat(wbp).st_mtime_ns
        prev = os.path.join(d, "preview")
        rc, out, err = run("claims-append", "--nas-root", nas, "--date", TODAY, "--preview-dir", prev)
        check(rc == 0, f"exit 0 (got {rc}; {err.strip()[-300:]})")
        check(sha(wbp) == before and os.stat(wbp).st_mtime_ns == mtime, "the workbook is untouched (bytes and mtime)")
        check(sorted(os.listdir(os.path.join(nas, "projects"))) == [cw.CLAIMS_NAME], "nothing written beside it")
        check(os.path.isfile(os.path.join(prev, "_MRI sessions - who ran them - preview.xlsx")), "the preview exists")
        check(os.path.isfile(os.path.join(prev, "claims_append_sessions.csv")), "and the report")
        check("to append: 2 sessions (3 acquisitions); 5 acquisition rows (2 of them in sessions already listed)" in out,
              "counts: 2 new sessions, 5 acquisition rows, 2 of them in the listed session")
        check("preview verified" in out, "the preview verifies")
        rc, _o, err = run("claims-append", "--nas-root", nas, "--preview-dir", os.path.join(nas, "projects", "pv"))
        check(rc == 1 and not os.path.exists(os.path.join(nas, "projects", "pv")), "a preview dir on the NAS: refused")


def test_claims_apply():
    print("claims-append: --apply appends exactly the new sessions and acquisitions:")
    with tempfile.TemporaryDirectory() as d:
        nas, wbp = setup_claims(d)
        orig = os.path.join(d, "orig.xlsx")
        with open(wbp, "rb") as f, open(orig, "wb") as g:
            g.write(f.read())
        bk = os.path.join(d, "bk")
        rc, out, err = run("claims-append", "--nas-root", nas, "--date", TODAY, "--apply", "--backup-dir", bk)
        check(rc == 0, f"exit 0 (got {rc}; {err.strip()[-300:]})")
        check(sha(os.path.join(bk, cw.CLAIMS_NAME)) == sha(orig), "the copy taken first is the old workbook")
        old, new = cells(orig), cells(wbp)
        check(all(new.get(k) == v for k, v in old.items()), "every existing cell keeps its value and number format")
        wb = load_workbook(wbp)
        ws = wb["Sessions to claim"]
        check(ws.max_row == 4, f"2 session rows appended (rows 3-4; max_row {ws.max_row})")
        rows = [[c.value for c in ws[r]] for r in (3, 4)]
        check(ws.cell(row=1, column=12).value == "Added" and ws.cell(row=2, column=12).value is None,
              "an 'Added' column after the last one, empty on the existing row")
        check(rows[0][0] == "(no project)" and rows[0][10] == "20210420_155137_jrc210420_m62_1019_1_1"
              and rows[0][11] == f"{TODAY}, Proyecto 1019 (2021) recovery",
              "the 1019 session: no project, keyed by its study folder, Added = date + source")
        check(rows[1][0] == "AE-biomaGUNE-0118" and rows[1][5] == 2 and rows[1][4] == "r111_0118"
              and rows[1][11] == f"{TODAY}, M. Jesús drive"
              and rows[1][9] == "ACQ-20190909-MRI-001 … ACQ-20190909-MRI-002",
              "the drive-3 session: project, 2 acquisitions (one config spelled with backslashes), ACQ-IDs as built")
        check(rows[1][1] == dt.datetime(2019, 9, 9) and ws.cell(row=4, column=2).number_format == "yyyy-mm-dd"
              and rows[1][2] == "12:42", "date and start from the folder name, in the sheet's date format")
        check(ws.cell(row=4, column=7).fill.fgColor.rgb == FILL.fgColor.rgb, "the new rows keep the yellow answer columns")
        check(ws.auto_filter.ref == "A1:L4", f"the autofilter covers the new rows and column ({ws.auto_filter.ref})")
        dvr = " ".join(str(dv.sqref) for dv in ws.data_validations.dataValidation)
        check("G3:G4" in dvr, f"the 'Who ran it' validation is extended ({dvr})")
        wa = wb["Acquisitions"]
        got = [wa.cell(row=r, column=1).value for r in range(4, wa.max_row + 1)]
        check(sorted(got) == ["ACQ-20190909-MRI-001", "ACQ-20190909-MRI-002", "ACQ-20210420-MRI-001",
                              "ACQ-20220124-MRI-003", "ACQ-20220124-MRI-004"], f"5 acquisition rows appended ({got})")
        r4 = [wa.cell(row=r, column=1).value for r in range(1, wa.max_row + 1)].index("ACQ-20220124-MRI-004") + 1
        check(wa.cell(row=r4, column=2).value == S1 and wa.cell(row=r4, column=3).value == "9",
              "a '<study>__9' name joins its study, exam 9")
        wr = wb["Read me"]
        line = wr.cell(row=wr.max_row, column=1).value or ""
        check(wr.max_row == 5 and line.startswith(f"Added {TODAY}: 2 sessions (3 acquisitions)")
              and "2 more acquisitions of sessions already listed" in line, f"one Read me line ({line[:90]}...)")
        check(wb["Projects"].max_row == 3, "the Projects summary is left as it is")
        before = sha(wbp)
        rc, out, _e = run("claims-append", "--nas-root", nas, "--date", TODAY, "--apply",
                          "--backup-dir", os.path.join(d, "bk2"))
        check(rc == 0 and "nothing to append" in out and sha(wbp) == before and not os.path.exists(os.path.join(d, "bk2")),
              "a re-run appends nothing, writes nothing, takes no copy")


def test_claims_refusals():
    print("claims-append: refusals:")
    with tempfile.TemporaryDirectory() as d:
        nas, wbp = setup_claims(d)
        before = sha(wbp)
        lock = os.path.join(nas, "projects", "~$" + cw.CLAIMS_NAME)
        with open(lock, "wb") as f:
            f.write(b"x")
        rc, _o, err = run("claims-append", "--nas-root", nas, "--apply", "--backup-dir", os.path.join(d, "bk"))
        check(rc == 1 and sha(wbp) == before and not os.path.exists(os.path.join(d, "bk")) and "open in Excel" in err,
              "Excel's lock file: refused, nothing written, no copy taken")
        rc, _o, _e = run("claims-append", "--nas-root", nas, "--preview-dir", os.path.join(d, "pv"))
        check(rc == 1, "the dry run refuses too")
        os.remove(lock)
        trunc = os.path.join(nas, "projects", "~$" + cw.CLAIMS_NAME[2:])
        with open(trunc, "wb") as f:
            f.write(b"x")
        rc, _o, _e = run("claims-append", "--nas-root", nas, "--preview-dir", os.path.join(d, "pv"))
        check(rc == 1, "a truncated owner-file name counts as a lock as well")
        os.remove(trunc)
        rc, _o, _e = run("claims-append", "--nas-root", nas, "--apply")
        check(rc == 2 and sha(wbp) == before, "--apply without --backup-dir: exit 2")
        os.makedirs(os.path.join(d, "exists"))
        rc, _o, err = run("claims-append", "--nas-root", nas, "--apply", "--backup-dir", os.path.join(d, "exists"))
        check(rc == 1 and sha(wbp) == before and "already exists" in err, "an existing --backup-dir: refused")
        rc, _o, err = run("claims-append", "--nas-root", nas, "--apply", "--backup-dir", os.path.join(nas, "bk"))
        check(rc == 1 and sha(wbp) == before and not os.path.exists(os.path.join(nas, "bk")), "a --backup-dir on the NAS: refused")
    with tempfile.TemporaryDirectory() as d:
        rows = claims_rows() + [reg_row("ACQ-20221024-CELL-001", instrument="CELL", operator=HOLD,
                                        original_name="drive3_MJesus-MFB/a/b.czi")]
        nas, wbp = setup_claims(d, rows)
        rc, _o, err = run("claims-append", "--nas-root", nas, "--preview-dir", os.path.join(d, "pv"))
        check(rc == 1 and "MRI only" in err, "a held non-MRI row: refused")
    with tempfile.TemporaryDirectory() as d:
        nas, wbp = setup_claims(d)
        wb = load_workbook(wbp)
        wb["Sessions to claim"].cell(row=1, column=11, value="session key")
        wb.save(wbp)
        before = sha(wbp)
        rc, _o, err = run("claims-append", "--nas-root", nas, "--apply", "--backup-dir", os.path.join(d, "bk"))
        check(rc == 1 and sha(wbp) == before and "headers edited" in err, "an edited header: refused")


def test_verification():
    print("compare_existing and the final check:")
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "a.xlsx")
        claims_workbook(p)
        a = load_workbook(p)
        check(cw.compare_existing(a, load_workbook(p), {}) == [], "an unchanged workbook verifies")
        for label, edit in (("a changed cell", lambda wb: wb["Sessions to claim"].cell(row=2, column=7, value="Marta")),
                            ("a cleared cell", lambda wb: setattr(wb["Acquisitions"].cell(row=3, column=6), "value", None)),
                            ("a changed number format",
                             lambda wb: setattr(wb["Sessions to claim"].cell(row=2, column=2), "number_format", "General")),
                            ("a value added to an existing row",
                             lambda wb: wb["Sessions to claim"].cell(row=2, column=13, value="x")),
                            ("an unplanned new row", lambda wb: wb["Projects"].cell(row=9, column=1, value="x"))):
            b = load_workbook(p)
            edit(b)
            check(cw.compare_existing(a, b, {}) != [], f"caught: {label}")
        b = load_workbook(p)
        b.create_sheet("extra")
        check(cw.compare_existing(a, b, {}) != [], "caught: a sheet added")

    with tempfile.TemporaryDirectory() as d:                 # a failed final check puts the copy back
        nas, wbp = setup_claims(d)
        before = sha(wbp)
        calls = []
        orig = cw.compare_existing

        def second_fails(*a, **k):
            calls.append(1)
            return orig(*a, **k) if len(calls) == 1 else ["simulated: a cell changed on disk"]

        with patched(cw, "compare_existing", second_fails):
            rc, _o, err = run("claims-append", "--nas-root", nas, "--apply", "--backup-dir", os.path.join(d, "bk"))
        check(rc == 1 and sha(wbp) == before and "putting the copy back" in err,
              "the final check fails: exit 1, the workbook byte-identical to before")
        leftovers = [f for f in os.listdir(os.path.join(nas, "projects")) if f.endswith(".tmp")]
        check(not leftovers, "no temp file is left beside the workbook")

        def refusing_replace(*_a, **_k):
            raise PermissionError("the file is open")

        with patched(cw.os, "replace", refusing_replace):
            rc, _o, err = run("claims-append", "--nas-root", nas, "--apply", "--backup-dir", os.path.join(d, "bk2"))
        check(rc == 1 and sha(wbp) == before and "unchanged" in err, "a swap that fails: exit 1, the workbook unchanged")
        check(not [f for f in os.listdir(os.path.join(nas, "projects")) if f.endswith(".tmp")], "and no temp file left")


# ---- 2. assign-append ----------------------------------------------------------------

MANIFEST_FIELDS = NP.MANIFEST_FIELDS
PILI = "Pili y Mili\\Draft papers"


def mrow(n, relpath, decision="holding", reason="no claim", verdict="", claim_id="", size=600_000_000):
    r = dict.fromkeys(MANIFEST_FIELDS, "")
    r.update({"row": str(n), "drive": "D3", "drive_label": "drive3_MJesus-MFB", "relpath": relpath, "size": str(size),
              "verdict": verdict, "claim_id": claim_id, "decision": decision, "reason": reason,
              "dest_rel": NP.holding_rel("MJesus-MFB", relpath)})
    return r


def assign_fixture(d):
    manifest = [
        mrow(1, PILI + "\\paper1\\fig1.tif"),
        mrow(2, PILI + "\\paper2\\Año 2.docx", size=500_000_000),
        mrow(3, "Otros\\stuff\\a.xlsx", reason="(C) claim", verdict="C", claim_id="CL-0957", size=1000),
        mrow(4, "Otros\\Slicer\\model.pth", reason="A2 D8: the group's own 3D Slicer tool"),        # not mappable
        mrow(5, "Microscopio\\x\\placed.tif", decision="place", reason="claim"),                   # not holding
    ]
    mpath = os.path.join(d, "placement_manifest.csv")
    with io.open(mpath, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=MANIFEST_FIELDS)
        w.writeheader()
        w.writerows(manifest)
    cell_rel = "Microscopio\\Microscopio- MJesus Sanchez 2023\\Machos vs Hembras\\WGA\\ID134-WGA20XLV-2.czi"
    exam = "MRI\\Proyecto 0619\\Flujo\\PA Flow Marzo 2020\\20191015_153156_jrc191015_m174_flow_1_1\\1"
    claims = [{"drive": "drive3", "relpath": cell_rel, "claim_id": "", "verdict": "NO-CLAIM", "proposed_project": ""},
              {"drive": "drive3", "relpath": exam + "\\pdata\\1\\dicom\\a.dcm", "claim_id": "CL-0011", "verdict": "C",
               "proposed_project": "0619"},
              {"drive": "drive3", "relpath": exam + "\\pdata\\1\\dicom\\b.dcm", "claim_id": "CL-0011", "verdict": "C",
               "proposed_project": "0619"}]
    cpath = os.path.join(d, "file_claims.csv")
    with io.open(cpath, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["drive", "relpath", "claim_id", "verdict", "proposed_project"])
        w.writeheader()
        w.writerows(claims)
    rows = [
        reg_row("ACQ-20221024-CELL-001", instrument="CELL", operator="MJ",
                original_name="drive3_MJesus-MFB/" + cell_rel.replace("\\", "/"),
                ingest_config="tools/configs/drives3_2026-10/drives3_C01.yaml"),
        reg_row("ACQ-20191015-MRI-001", notes=f"animal m174; historical drive 3 (M. Jesus, MJesus-MFB, WD WX22D623YP29): "
                                              f"{exam}; scanner-exported DICOM",
                original_name="20191015_153156_jrc191015_m174_flow_1_1/1",
                ingest_config="tools/configs/drive3_mri/drive3_mri_M08_noproject.yaml"),
        reg_row("ACQ-20221024-CELL-002", instrument="CELL", project_id="PROJ-0004",
                original_name="drive3_MJesus-MFB/Microscopio/y.czi",
                ingest_config="tools/configs/drives3_2026-10/drives3_C02.yaml"),               # has a project
        reg_row("ACQ-20240304-CELL-001", instrument="CELL", original_name="drive1_FRIO-X6/a.czi",
                ingest_config="tools/configs/drives_2026-09/x.yaml"),                          # another drive
    ]
    nas = write_nas(d, rows, PROJECTS)
    wbp = os.path.join(nas, "projects", cw.ASSIGN_NAME)
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    ws.cell(row=1, column=1, value="Historical drives: which project does this belong to?")
    wg = wb.create_sheet("Groups to assign")
    gcols = ["Group", "Priority", "Project", "Notes", "Answered by", "Researcher folder", "Drive", "Folder on the drive",
             "Raw acquisitions", "Other files", "Other files (GB)", "From", "To", "Example file names",
             "Where the other files are now", "Raw acquisitions (first .. last ACQ-ID)", "group key (do not edit)"]
    for j, n in enumerate(gcols, 1):
        wg.cell(row=1, column=j, value=n)
    for i, (g, key) in enumerate((("G001", "D2|Laura|Proyectos_Laboratorio_Laura\\Lung-surfactant"),
                                  ("G290", "D1|Claudia|Cell observer\\AINHIZE")), 2):
        wg.cell(row=i, column=1, value=g)
        wg.cell(row=i, column=3, value="laura" if g == "G001" else None).fill = FILL
        wg.cell(row=i, column=17, value=key)
    wg.auto_filter.ref = "A1:Q3"
    dv = DataValidation(type="list", formula1="=Projects!$A$2:$A$3", allow_blank=True)
    wg.add_data_validation(dv)
    dv.add("C2:C3")
    wa = wb.create_sheet("Acquisitions")
    for j, n in enumerate(["Group", "ACQ-ID", "Acquired", "Instrument", "File name", "Folder on the drive", "Drive",
                           "Researcher folder", "Claimed project (uncertain)"], 1):
        wa.cell(row=1, column=j, value=n)
    for j, v in enumerate(["G290", "ACQ-20240904-CELL-001", "2024-09-04", "CELL", "a.czi", "Cell observer\\AINHIZE",
                           "FRIO X6", "Claudia", None], 1):
        wa.cell(row=2, column=j, value=v)
    wp = wb.create_sheet("Projects")
    wp.cell(row=1, column=1, value="Project name")
    wp.cell(row=2, column=1, value="AE-biomaGUNE-0619")
    wb.save(wbp)
    return nas, wbp, mpath, cpath, manifest


def test_assign_apply():
    print("assign-append: drive-3 material grouped by the 2b functions, appended after the last G:")
    with tempfile.TemporaryDirectory() as d:
        nas, wbp, mpath, cpath, manifest = assign_fixture(d)
        orig = os.path.join(d, "orig.xlsx")
        with open(wbp, "rb") as f, open(orig, "wb") as g:
            g.write(f.read())
        prev = os.path.join(d, "pv")
        rc, out, err = run("assign-append", "--nas-root", nas, "--manifest", mpath, "--claims", cpath,
                           "--date", TODAY, "--preview-dir", prev)
        check(rc == 0 and sha(wbp) == sha(orig), f"dry run: exit 0, the workbook untouched (rc {rc}; {err.strip()[-200:]})")
        check("drive-3 holding files that 2b can map: 3; acquisitions with a blank project: 2" in out,
              "3 mappable holding files (not the D8 tool, not the placed row) and 2 blank-project acquisitions")
        bk = os.path.join(d, "bk")
        rc, out, err = run("assign-append", "--nas-root", nas, "--manifest", mpath, "--claims", cpath,
                           "--date", TODAY, "--apply", "--backup-dir", bk)
        check(rc == 0, f"--apply: exit 0 (got {rc}; {err.strip()[-300:]})")
        old, new = cells(orig), cells(wbp)
        check(all(new.get(k) == v for k, v in old.items()), "every existing cell unchanged (the answer 'laura' too)")
        wb = load_workbook(wbp)
        wg = wb["Groups to assign"]
        h = {c.value: c.column for c in wg[1]}
        groups = {wg.cell(row=r, column=h["group key (do not edit)"]).value: [c.value for c in wg[r]]
                  for r in range(4, wg.max_row + 1)}
        pili_key = NP.key_string(*NP.manifest_group(manifest[0])[:3])
        c_key = NP.key_string(*NP.manifest_group(manifest[2])[:3])
        check(pili_key == "D3||" + PILI and c_key == "D3||(C) CL-0957", f"(the 2b keys: {pili_key!r}, {c_key!r})")
        check(len(groups) == 4 and pili_key in groups and c_key in groups,
              f"4 new groups: the two holding groups and the two acquisitions' groups ({sorted(groups)})")
        pili = groups[pili_key]
        check(pili[h["Group"] - 1] == "G291" and pili[h["Priority"] - 1] == "A" and pili[h["Other files"] - 1] == 2
              and pili[h["Other files (GB)"] - 1] == 1.1, "the 1.1 GB group is G291, priority A, 2 files")
        check(pili[h["Where the other files are now"] - 1] ==
              "\\\\GJESUS3\\gjesus3\\gjesus3-data\\staging\\historical_drives_unassigned\\MJesus-MFB\\Pili y Mili\\Draft papers",
              "its files are shown in the holding folder, with the gjesus3-data share path")
        check(pili[h["Added"] - 1] == f"{TODAY}, M. Jesús drive" and pili[h["Project"] - 1] is None,
              "Added = date + source; Project left for the answer")
        mri_key = NP.key_string("D3", "", "(C) CL-0011")
        cell_key = NP.key_string("D3", "", "Microscopio\\Microscopio- MJesus Sanchez 2023")
        check(mri_key in groups and cell_key in groups,
              "the MRI exam takes the claim its files share ((C) CL-0011); the Cell Observer file its folder series")
        check(sorted(g[0] for g in groups.values()) == ["G291", "G292", "G293", "G294"], "numbered G291-G294")
        wa = wb["Acquisitions"]
        acq = {wa.cell(row=r, column=2).value: [c.value for c in wa[r]] for r in range(3, wa.max_row + 1)}
        check(set(acq) == {"ACQ-20221024-CELL-001", "ACQ-20191015-MRI-001"}, "2 acquisition rows appended")
        check(acq["ACQ-20191015-MRI-001"][0] == groups[mri_key][0] and acq["ACQ-20191015-MRI-001"][8] == "0619",
              "an acquisition row names its group's G-number, and the claimed project of a (C) claim")
        dvr = " ".join(str(dv.sqref) for dv in wg.data_validations.dataValidation)
        check("C4:C7" in dvr and wg.auto_filter.ref == "A1:R7", f"drop-down and filter extended ({dvr}; {wg.auto_filter.ref})")
        comp = os.path.join(bk, "drive3_blank_project_list.csv")
        lst = cw.read_csv(comp)
        check([r for r in lst] and set(lst[0]) == set(cw.BLANK_LIST_FIELDS), "the companion list has the blank-list columns")
        check({NP.key_string(*NP.blank_list_group(a)[:3]) for a in lst} == {mri_key, cell_key},
              "re-grouped by blank_list_group, the companion list gives the keys written in the workbook")

        # The point of it all: an answer typed against a group applies with the 2b tools.
        mapping = {NP.norm_key(pili_key): "AE-biomaGUNE-0619"}
        projects = {p["name"]: p for p in PROJECTS}
        out_rows, used = NP.remap_rows(cw.read_csv(mpath), mapping, projects)
        redecided = [r for r in out_rows if r["reason"].startswith("2b mapping")]
        check(len(redecided) == 2 and all(r["decision"] == "place" and r["project_name"] == "AE-biomaGUNE-0619"
                                          for r in redecided),
              "nonraw_placement.remap_rows applies an answer on that group key to exactly its 2 holding files")

        before = sha(wbp)
        rc, out, _e = run("assign-append", "--nas-root", nas, "--manifest", mpath, "--claims", cpath,
                          "--date", TODAY, "--apply", "--backup-dir", os.path.join(d, "bk2"))
        check(rc == 0 and "nothing to append" in out and sha(wbp) == before, "a re-run appends nothing")


def test_assign_refuses_an_unplaceable_acquisition():
    print("assign-append: an acquisition with no drive path is refused, not guessed:")
    with tempfile.TemporaryDirectory() as d:
        nas, wbp, mpath, cpath, _m = assign_fixture(d)
        path = os.path.join(nas, "registries", "registry_raw.csv")
        with open(path, "a", encoding="utf-8", newline="") as f:
            r = reg_row("ACQ-20191016-MRI-001", original_name="20191016_1_1/1", notes="no path here",
                        ingest_config="tools/configs/drive3_mri/drive3_mri_M08_noproject.yaml")
            csv.writer(f, lineterminator="\r\n").writerow([r[k] for k in registry.REGISTRY_FIELDS])
        before = sha(wbp)
        rc, _o, err = run("assign-append", "--nas-root", nas, "--manifest", mpath, "--claims", cpath,
                          "--preview-dir", os.path.join(d, "pv"))
        check(rc == 1 and sha(wbp) == before and "refusing to guess" in err, "refused, nothing written")


def test_session_key():
    print("session_key: the builder's key, and the `__<exam>` names:")
    for name, want in ((S1 + "/12", S1), (S1 + "__4", S1), ("a/b/c", "a/b"), ("noslash", None), ("__4", None)):
        check(cw.session_key(name) == want, f"{name!r} -> {want!r}")


def main():
    for fn in (test_session_key, test_claims_dry_run, test_claims_apply, test_claims_refusals, test_verification,
               test_assign_apply, test_assign_refuses_an_unplaceable_acquisition):
        fn()
    print()
    if FAILS:
        print(f"FAILED ({len(FAILS)}):")
        for m in FAILS:
            print(f"  - {m}")
        sys.exit(1)
    print("ALL CHECKS PASSED")


if __name__ == "__main__":
    main()
