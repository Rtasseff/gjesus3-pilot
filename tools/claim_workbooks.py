#!/usr/bin/env python3
"""claim_workbooks.py -- APPEND to the two shared workbooks in projects\\, never rebuild them.
Data Office only. Dry run by default.

    python tools/claim_workbooks.py claims-append --nas-root "J:\\gjesus3-data"                  (dry run)
    python tools/claim_workbooks.py claims-append --nas-root "J:\\gjesus3-data" --apply \\
        --backup-dir "C:\\Users\\rtasseff\\temp\\gjesus3_claims_append_YYYYMMDD"
    python tools/claim_workbooks.py assign-append --nas-root "J:\\gjesus3-data"                  (dry run)
    python tools/claim_workbooks.py assign-append --nas-root "J:\\gjesus3-data" --apply --backup-dir ...

WHY. Ryan, 2026-10-07 (STATUS section 0.5 "Hold"): historical internal MRI loaded without an
operator gets `operator = pending-claim`, and its sessions are listed in the claim workbook,
which people fill in; the two workbooks are appended to, never rebuilt (people's answers live
in them). 06_REGISTRIES section 2.3a-bis.

claims-append  -> projects\\_MRI sessions - who ran them.xlsx
  One row per MRI SESSION whose acquisitions hold `operator = pending-claim` and that the sheet
  "Sessions to claim" does not list yet, keyed exactly as the workbook was built: the session
  key is the part of `original_name` before the last `/` (the scanner study folder). A name
  with no `/` but a `__<exam>` suffix (3 drives 1+2 rows, `<study>__4`) joins its study. Each
  acquisition not yet on the sheet "Acquisitions" is appended there too, so a new acquisition
  of a session already listed is shown (the session's answer covers it). The sheet "Projects"
  (a summary with totals) is left as it is.

assign-append  -> projects\\_Historical data - assign to projects.xlsx
  The M. Jesus drive's material with no project, grouped EXACTLY as the 2b mapping groups it
  (tools/drive_staging/nonraw_placement.py manifest_group / blank_list_group / key_string), so
  an answer typed against a group applies with the existing 2b tools
  (tasks/drives_nonraw_2b_2c_followup.md: `remap --manifest <drive-3 manifest>`, `apply-raw`):
    - its holding-folder files (staging\\historical_drives_unassigned\\MJesus-MFB\\): the rows of
      the drive-3 placement manifest with decision `holding` and a mappable reason;
    - its acquisitions registered with a blank project (default: the Cell Observer configs
      tools/configs/drives3_2026-10/ and stream M's tools/configs/drive3_mri/). Their drive path
      is the `original_name` below `drive3_MJesus-MFB/` (Cell Observer) or the path in `notes`
      after "WD WX22D623YP29): " (MRI); their claim comes from part A2's file_claims.csv.
  New groups are numbered after the highest G-number on the sheet. A companion
  `drive3_blank_project_list.csv` (the columns of tasks/drives_blank_project_list.csv, which
  `apply-raw` reads) is written beside the outputs for the raw half.

BOTH. Append only: every existing cell keeps its value and number format (checked after
writing, against a copy taken first), new rows go below the last row, the new column `Added`
(date + source) is added after the last column and left empty on the existing rows, and one
line goes at the end of "Read me". The autofilter and the drop-down / text-length validation
are extended over the new rows. Refuses when Excel's lock file `~$<name>` is beside the
workbook (someone has it open). A dry run writes a preview copy (and a report) to
--preview-dir, never beside the workbook. --apply needs --backup-dir (a folder that does not
exist yet, off the NAS): the workbook is copied there and SHA-256-checked first; the new file
is built and verified there, then copied in beside the workbook and swapped in with
os.replace, then verified again; on a mismatch the copy is put back.

Exit codes: 0 done / dry run / nothing to append; 1 refused, or verification failed and the
copy was restored; 2 bad arguments; 3 verification failed AND the restore failed.
"""

import argparse
import collections
import csv
import datetime as dt
import hashlib
import io
import os
import re
import shutil
import sys
import tempfile
from copy import copy

_THIS_DIR = os.path.dirname(os.path.abspath(__file__))
for _p in (_THIS_DIR, os.path.join(_THIS_DIR, "drive_staging")):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from ingest import registry  # noqa: E402

HOLD = registry.OPERATOR_HOLD
CLAIMS_NAME = "_MRI sessions - who ran them.xlsx"
ASSIGN_NAME = "_Historical data - assign to projects.xlsx"
ADDED = "Added"
DATA_SHARE = "\\\\GJESUS3\\gjesus3\\gjesus3-data"      # J:\gjesus3-data as other machines see it

# The builder's study-folder form, YYYYMMDD_HHMMSS_...
STUDY_RE = re.compile(r"^(\d{4})(\d{2})(\d{2})_(\d{2})(\d{2})(\d{2})_")

# Where a held MRI session came from, by its ingest_config (slashes normalised); for the Added cell.
SOURCES = (
    ("tools/configs/drive3_mri/", "M. Jesús drive"),
    ("tools/configs/drives_2026-09/", "FRIO X6 / MFB Disco 2 drives"),
    ("tools/configs/mri_1019_kgjesus_2021_", "Proyecto 1019 (2021) recovery"),
    ("tools/configs/mri_0522_m145", "MRI scanner, 2025 session m145 (0522)"),
    ("tools/configs/mri_july_1125/", "MRI scanner, 2026 phantom studies"),
)

# The M. Jesus drive (drive 3)
D3 = "D3"
D3_DISPLAY = "MJesus-MFB"
D3_NAME_PREFIX = "drive3_MJesus-MFB/"
D3_NOTES_RE = re.compile(r"WD WX22D623YP29\): ([^;]+);")
D3_MANIFEST = r"D:\projects\gjesus3\drive3_streams\placement\manifest_v2\placement_manifest.csv"
D3_CLAIMS = r"D:\projects\gjesus3\drive3_analysis\a2\claims\file_claims.csv"
D3_PREFIXES = ("tools/configs/drives3_2026-10/", "tools/configs/drive3_mri/")
D3_SOURCE = "M. Jesús drive"
BLANK_LIST_FIELDS = ["acq_id", "batch", "instrument", "acquisition_date", "verdict", "path_claimed_project",
                     "claim_id", "researcher", "drive", "relpath", "archive", "member", "other_copies"]


class WorkbookError(Exception):
    """A refusal: the run stops; unless it says otherwise, nothing was written."""


def log(msg, level="INFO"):
    print(f"[{dt.datetime.now():%H:%M:%S}] {level}: {msg}", file=sys.stderr)


# ---- Small helpers ------------------------------------------------------------

def norm_config(value):
    return (value or "").replace("\\", "/")


def sha256_bytes(data):
    return hashlib.sha256(data).hexdigest()


def sha256_file(path):
    with open(path, "rb") as f:
        return sha256_bytes(f.read())


def read_bytes(path):
    with open(path, "rb") as f:
        return f.read()


def read_csv(path):
    with io.open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_csv(path, fields, rows):
    with io.open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def is_inside(child, parent):
    if not parent:
        return False
    child = os.path.normcase(os.path.abspath(child))
    parent = os.path.normcase(os.path.abspath(parent))
    try:
        return os.path.commonpath([child, parent]) == parent
    except ValueError:                                  # different drives
        return False


def lock_files(path):
    """Excel's owner files for `path` that exist (Excel writes `~$` + the name; Word-style
    truncations of the first one or two characters are checked as well)."""
    folder, name = os.path.split(os.path.abspath(path))
    seen = []
    for cand in ("~$" + name, "~$" + name[1:], "~$" + name[2:]):
        p = os.path.join(folder, cand)
        if p not in seen and os.path.exists(p):
            seen.append(p)
    return seen


def refuse_if_locked(path):
    locks = lock_files(path)
    if locks:
        raise WorkbookError(f"{os.path.basename(path)} is open in Excel (lock file {os.path.basename(locks[0])}). "
                            f"Ask whoever has it open to close it, then re-run. Nothing was written.")


def headers(ws):
    """{header text: 1-based column} of row 1."""
    return {c.value: c.column for c in ws[1] if c.value is not None}


def need(ws, names):
    h = headers(ws)
    missing = [n for n in names if n not in h]
    if missing:
        raise WorkbookError(f"sheet {ws.title!r} has no column {missing[0]!r} (were the headers edited?): "
                            f"refusing to append.")
    return h


def sheet(wb, title):
    if title not in wb.sheetnames:
        raise WorkbookError(f"the workbook has no sheet {title!r}: refusing to append.")
    return wb[title]


def last_row(ws):
    """The last row holding any value (ignores trailing rows that carry only formatting)."""
    for r in range(ws.max_row, 0, -1):
        if any(c.value is not None for c in ws[r]):
            return r
    return 0


def last_col(ws):
    """The last column of the header row."""
    cols = [c.column for c in ws[1] if c.value is not None]
    return max(cols) if cols else 0


def column_values(ws, col, first=2):
    return [ws.cell(row=r, column=col).value for r in range(first, last_row(ws) + 1)]


# ---- Appending to one sheet -----------------------------------------------------

class SheetAppend:
    """Rows to append to one sheet: values by header name (None = empty)."""

    def __init__(self, title, rows, added_column=True):
        self.title = title
        self.rows = rows                     # [{header: value}]
        self.added_column = added_column


def apply_appends(wb, appends, readme_line):
    """Append in memory. Returns {sheet title: (old last row, number of new rows, expected new last row)}."""
    done = {}
    for ap in appends:
        if not ap.rows:
            continue
        ws = sheet(wb, ap.title)
        h = headers(ws)
        old_last, old_lastcol = last_row(ws), last_col(ws)
        added_col = None
        if ap.added_column:
            added_col = h.get(ADDED)
            if added_col is None:
                added_col = old_lastcol + 1
                src = ws.cell(row=1, column=old_lastcol)
                c = ws.cell(row=1, column=added_col, value=ADDED)
                c._style = copy(src._style)
                ws.column_dimensions[get_column_letter(added_col)].width = 30
                h[ADDED] = added_col
        for name in ap.rows[0]:
            if name not in h:
                raise WorkbookError(f"sheet {ap.title!r} has no column {name!r}: refusing to append.")
        template_row = old_last if old_last >= 2 else None
        for i, values in enumerate(ap.rows, start=old_last + 1):
            if template_row is not None:                 # the look of the last existing row, every column
                for col in range(1, old_lastcol + 1):
                    ws.cell(row=i, column=col)._style = copy(ws.cell(row=template_row, column=col)._style)
            for name, value in values.items():
                ws.cell(row=i, column=h[name]).value = value
        new_last = old_last + len(ap.rows)
        if ws.auto_filter.ref:
            ws.auto_filter.ref = f"A1:{get_column_letter(max(last_col(ws), old_lastcol))}{new_last}"
        for dv in ws.data_validations.dataValidation:
            for rng in list(dv.sqref.ranges):
                if rng.max_row == old_last and rng.min_row <= 2:
                    dv.add(f"{get_column_letter(rng.min_col)}{old_last + 1}:{get_column_letter(rng.max_col)}{new_last}")
        done[ap.title] = (old_last, len(ap.rows), new_last)
    if readme_line and done:
        ws = sheet(wb, "Read me")
        old_last = last_row(ws)
        r = old_last + 2                                 # one empty row, then the line
        src = ws.cell(row=max(1, old_last), column=1)
        c = ws.cell(row=r, column=1, value=readme_line)
        c._style = copy(src._style)
        done["Read me"] = (old_last, 1, r)
    return done


def compare_existing(orig_wb, new_wb, done):
    """Problems (empty = verified): every existing cell of every sheet unchanged (value, type and
    number format), no sheet added, removed or renamed, nothing new in an existing row except the
    `Added` header, and each sheet's last row exactly where the plan put it."""
    problems = []
    if orig_wb.sheetnames != new_wb.sheetnames:
        return [f"the sheets changed: {orig_wb.sheetnames} -> {new_wb.sheetnames}"]
    for title in orig_wb.sheetnames:
        a, b = orig_wb[title], new_wb[title]
        a_last, a_cols = a.max_row, a.max_column
        for r in range(1, a_last + 1):
            for c in range(1, a_cols + 1):
                x, y = a.cell(row=r, column=c), b.cell(row=r, column=c)
                if x.value != y.value or type(x.value) is not type(y.value) or x.number_format != y.number_format:
                    problems.append(f"{title}!{get_column_letter(c)}{r}: {x.value!r} -> {y.value!r}")
                    if len(problems) >= 10:
                        return problems + ["... (more)"]
            for c in range(a_cols + 1, b.max_column + 1):
                v = b.cell(row=r, column=c).value
                if v is not None and not (r == 1 and v == ADDED):
                    problems.append(f"{title}!{get_column_letter(c)}{r}: a new value {v!r} in an existing row")
        want_last = done[title][2] if title in done else last_row(a)
        if last_row(b) != want_last:
            problems.append(f"{title}: the last row is {last_row(b)}, expected {want_last}")
    return problems


# ---- Writing ---------------------------------------------------------------------

def run_write(kind, workbook, data, wb, appends, readme_line, report, args):
    """Dry run (a preview copy) or --apply. `data` are the workbook bytes the plan was made from.
    Returns the exit code."""
    name = os.path.basename(workbook)
    orig_wb = load_workbook(io.BytesIO(data))
    done = apply_appends(wb, appends, readme_line)
    n_rows = {t: n for t, (_o, n, _c) in done.items()}
    if not done:
        print("nothing to append.")
        return 0

    if not args.apply:
        out = os.path.abspath(args.preview_dir)
        if is_inside(out, args.nas_root) or os.path.normcase(out) == os.path.normcase(os.path.dirname(os.path.abspath(workbook))):
            raise WorkbookError(f"--preview-dir {out} is on the NAS / beside the workbook; give a scratch folder.")
        os.makedirs(out, exist_ok=True)
        preview = os.path.join(out, f"{os.path.splitext(name)[0]} - preview.xlsx")
        wb.save(preview)
        problems = compare_existing(orig_wb, load_workbook(preview), done)
        for fn, (fields, rows) in report.items():
            write_csv(os.path.join(out, fn), fields, rows)
        print(f"DRY RUN - the workbook is not touched. Preview: {preview}")
        print("rows to append: " + ", ".join(f"{t}={n}" for t, n in n_rows.items()))
        if problems:
            for p in problems:
                log(p, "ERROR")
            raise WorkbookError("the preview does not verify (see above): the tool would refuse to write.")
        print("preview verified: every existing cell unchanged.")
        return 0

    backup_dir = os.path.abspath(args.backup_dir)
    if os.path.exists(backup_dir):
        raise WorkbookError(f"--backup-dir {backup_dir} already exists; give a folder that does not exist yet.")
    if is_inside(backup_dir, args.nas_root) or is_inside(backup_dir, os.path.dirname(os.path.abspath(workbook))):
        raise WorkbookError(f"--backup-dir {backup_dir} is on the NAS; the copy must be off the NAS.")
    refuse_if_locked(workbook)
    os.makedirs(backup_dir)
    backup = os.path.join(backup_dir, name)
    shutil.copy2(workbook, backup)
    sha0 = sha256_bytes(data)
    if not (sha256_file(backup) == sha256_file(workbook) == sha0):
        raise WorkbookError(f"the workbook changed since it was read, or its copy does not verify; nothing was "
                            f"written. See {backup_dir}.")
    log(f"copy taken: {backup} (SHA-256 verified)")
    new = os.path.join(backup_dir, f"{os.path.splitext(name)[0]} - new.xlsx")
    wb.save(new)
    problems = compare_existing(load_workbook(backup), load_workbook(new), done)
    if problems:
        for p in problems:
            log(p, "ERROR")
        raise WorkbookError("the new workbook does not verify against the copy; the workbook was not touched.")
    for fn, (fields, rows) in report.items():
        write_csv(os.path.join(backup_dir, fn), fields, rows)

    refuse_if_locked(workbook)                         # someone may have opened it meanwhile
    if sha256_file(workbook) != sha0:
        raise WorkbookError("the workbook changed while the new one was built; nothing was written.")
    tmp = os.path.join(os.path.dirname(os.path.abspath(workbook)), f".claim_workbooks.{os.getpid()}.tmp")
    try:
        shutil.copyfile(new, tmp)
        os.replace(tmp, workbook)
    except OSError as exc:
        if os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass
        unchanged = sha256_file(workbook) == sha0
        raise WorkbookError(f"could not put the new workbook in place ({exc}); the workbook is "
                            f"{'unchanged' if unchanged else 'CHANGED: compare it with ' + backup}.")
    problems = []
    if sha256_file(workbook) != sha256_file(new):
        problems.append("the workbook on disk is not the verified new file")
    else:
        problems = compare_existing(load_workbook(backup), load_workbook(workbook), done)
    if problems:
        for p in problems:
            log(p, "ERROR")
        log("verification FAILED: putting the copy back.", "ERROR")
        try:
            shutil.copyfile(backup, tmp)
            os.replace(tmp, workbook)
            restored = sha256_file(workbook) == sha0
        except OSError as exc:
            log(f"the restore failed ({exc})", "ERROR")
            restored = False
        if restored:
            log("the copy is back; the workbook is byte-identical to before.", "ERROR")
            return 1
        log(f"RESTORE FAILED: copy {backup} over {workbook} by hand.", "ERROR")
        return 3
    print()
    print(f"APPENDED ({kind})")
    for t, n in n_rows.items():
        print(f"rows_appended[{t}]={n}")
    print(f"sha256_before={sha0}")
    print(f"sha256_after={sha256_file(workbook)}")
    print(f"backup={backup}")
    return 0


def load_target(path):
    if not os.path.isfile(path):
        raise WorkbookError(f"workbook not found: {path}")
    refuse_if_locked(path)
    data = read_bytes(path)
    return data, load_workbook(io.BytesIO(data))


def read_registry(nas_root):
    path = os.path.join(nas_root, "registries", "registry_raw.csv")
    if not os.path.isfile(path):
        raise WorkbookError(f"registry_raw.csv not found under {nas_root}\\registries")
    with io.open(path, encoding="utf-8", newline="") as f:
        rows = list(csv.DictReader(f))
    projects = {p["project_id"]: p for p in read_csv(os.path.join(nas_root, "registries", "registry_projects.csv"))}
    return rows, projects


# ---- claims-append ----------------------------------------------------------------

def session_key(original_name):
    """The workbook's session key: the scanner study folder = `original_name` before the last `/`
    (as the builder keyed it). A name without `/` but with a `__<exam>` suffix joins its study."""
    study, sep, _exam = original_name.rpartition("/")
    if sep:
        return study
    study, sep, exam = original_name.rpartition("__")
    if sep and study and exam:
        return study
    return None


def exam_of(original_name):
    if "/" in original_name:
        return original_name.rpartition("/")[2]
    return original_name.rpartition("__")[2]


def source_of(ingest_config):
    cfg = norm_config(ingest_config)
    for prefix, label in SOURCES:
        if cfg.startswith(prefix):
            return label
    folder = os.path.basename(os.path.dirname(cfg)) or cfg
    return f"ingest config {folder}"


def ids_text(ids):
    ids = sorted(ids)
    if len(ids) == 1:
        return ids[0]
    if len(ids) == 2:
        return f"{ids[0]} … {ids[-1]}"
    return f"{ids[0]} … {ids[-1]} ({len(ids)})"


def plan_claims(rows, projects, wb, today):
    """(appends, readme line, report) for claims-append."""
    ws_s, ws_a = sheet(wb, "Sessions to claim"), sheet(wb, "Acquisitions")
    hs = need(ws_s, ["Project", "Date", "Start", "Session folder on the scanner", "Animal", "Acquisitions",
                     "ACQ-IDs (first … last)", "session key (do not edit)"])
    ha = need(ws_a, ["ACQ-ID", "Session folder on the scanner", "Exam", "Acquired", "Project", "Animal",
                     "Subject (facility id)"])
    have_sessions = {v for v in column_values(ws_s, hs["session key (do not edit)"]) if v}
    have_acqs = {v for v in column_values(ws_a, ha["ACQ-ID"]) if v}

    hit = [r for r in rows if r["operator"] == HOLD]
    other = sorted({r["instrument"] for r in hit} - {"MRI"})
    if other:
        raise WorkbookError(f"rows of {other} hold {HOLD!r}: the claim workbook is for MRI only; refusing.")
    by = collections.defaultdict(list)
    for r in hit:
        k = session_key(r["original_name"])
        if not k:
            raise WorkbookError(f"{r['acq_id']}: original_name {r['original_name']!r} names no study folder; "
                                f"refusing to guess its session.")
        by[k].append(r)

    def project_text(rs):
        names = []
        for pid in sorted({x["project_id"] for x in rs}):
            p = projects.get(pid, {})
            names.append(p.get("name", pid or "(no project)") + (" (closed)" if p.get("status") == "closed" else ""))
        return "; ".join(names)

    sessions = []
    for key, rs in by.items():
        if key in have_sessions:
            continue
        m = STUDY_RE.match(key)
        if m:
            y, mo, d, hh, mi, _s = (int(x) for x in m.groups())
            date, start = dt.datetime(y, mo, d), f"{hh:02d}:{mi:02d}"
        else:
            first = min(x["acquisition_datetime"] for x in rs)
            date, start = dt.datetime.strptime(first[:10], "%Y-%m-%d"), first[11:16]
        sources = sorted({source_of(x["ingest_config"]) for x in rs})
        sessions.append({
            "key": key, "date": date, "start": start, "project": project_text(rs),
            "animal": "; ".join(sorted({x["sample_id"] for x in rs if x["sample_id"]})),
            "n": len(rs), "ids": ids_text(x["acq_id"] for x in rs), "added": f"{today}, {' + '.join(sources)}",
            "sources": sources})
    sessions.sort(key=lambda s: (s["project"].lower(), s["date"], s["start"], s["key"]))
    s_rows = [{"Project": s["project"], "Date": s["date"], "Start": s["start"],
               "Session folder on the scanner": s["key"], "Animal": s["animal"] or None, "Acquisitions": s["n"],
               "ACQ-IDs (first … last)": s["ids"], "session key (do not edit)": s["key"], ADDED: s["added"]}
              for s in sessions]

    new_acqs = sorted((r for r in hit if r["acq_id"] not in have_acqs),
                      key=lambda r: (project_text([r]).lower(), session_key(r["original_name"]), r["acq_id"]))
    a_rows = [{"ACQ-ID": r["acq_id"], "Session folder on the scanner": session_key(r["original_name"]),
               "Exam": exam_of(r["original_name"]),
               "Acquired": r["acquisition_datetime"][:19].replace("T", " "), "Project": project_text([r]),
               "Animal": r["sample_id"] or None, "Subject (facility id)": r["subject_ids"] or None,
               ADDED: f"{today}, {source_of(r['ingest_config'])}"} for r in new_acqs]

    in_listed = sum(1 for r in new_acqs if session_key(r["original_name"]) in have_sessions)
    per_source = collections.Counter()
    for s in sessions:
        per_source[" + ".join(s["sources"])] += 1
    line = None
    if sessions or a_rows:
        parts = ", ".join(f"{n:,} from the {src}" for src, n in sorted(per_source.items()))
        line = (f"Added {today}: {len(sessions):,} sessions ({sum(s['n'] for s in sessions):,} acquisitions)"
                + (f": {parts}" if parts else "")
                + (f"; and {in_listed:,} more acquisitions of sessions already listed" if in_listed else "")
                + ". They are marked in the column 'Added' of 'Sessions to claim' and 'Acquisitions'. "
                  "They come from historical copies, where no record says who ran the scanner. "
                  "The sheet 'Projects' counts the first list only.")
    report = {"claims_append_sessions.csv": (
        ["session_key", "project", "date", "start", "acquisitions", "added"],
        [{"session_key": s["key"], "project": s["project"], "date": s["date"].date().isoformat(),
          "start": s["start"], "acquisitions": s["n"], "added": s["added"]} for s in sessions])}
    print(f"held MRI acquisitions (operator = {HOLD}): {len(hit):,} in {len(by):,} sessions")
    print(f"already in the workbook: {len(have_sessions):,} sessions, {len(have_acqs):,} acquisitions")
    print(f"to append: {len(sessions):,} sessions ({sum(s['n'] for s in sessions):,} acquisitions); "
          f"{len(a_rows):,} acquisition rows ({in_listed:,} of them in sessions already listed)")
    for src, n in sorted(per_source.items()):
        print(f"  {n:6,d} sessions from: {src}")
    return [SheetAppend("Sessions to claim", s_rows), SheetAppend("Acquisitions", a_rows)], line, report


def cmd_claims(args):
    workbook = args.workbook or os.path.join(args.nas_root, "projects", CLAIMS_NAME)
    rows, projects = read_registry(args.nas_root)
    data, wb = load_target(workbook)
    appends, line, report = plan_claims(rows, projects, wb, args.date)
    return run_write("claims-append", workbook, data, wb, appends, line, report, args)


# ---- assign-append ----------------------------------------------------------------

def d3_relpath(r):
    """The drive-3 path (below the drive root, backslashes) of a registered acquisition, or None."""
    name = r["original_name"]
    if name.startswith(D3_NAME_PREFIX):
        return name[len(D3_NAME_PREFIX):].replace("/", "\\")
    m = D3_NOTES_RE.search(r.get("notes", ""))
    return m.group(1).strip() if m else None


def claims_for(paths, claims_path):
    """{relpath: (verdict, claim_id, proposed_project)} from part A2's file_claims.csv: the file's own
    row, or for a folder (an MRI exam) the one claim every file under it shares ('' when they differ)."""
    want = set(paths)
    exact, under = {}, collections.defaultdict(set)
    for c in read_csv(claims_path):
        rel = c["relpath"]
        val = (c.get("verdict", ""), c.get("claim_id", ""), c.get("proposed_project", ""))
        if rel in want:
            exact[rel] = val
        parts = rel.split("\\")
        for i in range(1, len(parts)):
            folder = "\\".join(parts[:i])
            if folder in want:
                under[folder].add(val)
    out = {}
    for p in want:
        if p in exact:
            out[p] = exact[p]
        elif len(under.get(p, ())) == 1:
            out[p] = next(iter(under[p]))
        else:
            out[p] = ("", "", "")
    return out


def blank_project_acqs(rows, prefixes, claims_path):
    """The blank-list-shaped rows (tasks/drives_blank_project_list.csv columns) of the drive-3
    acquisitions registered with a blank project."""
    pref = [norm_config(p) for p in prefixes]
    sel = [r for r in rows if r["project_id"] == "" and any(norm_config(r["ingest_config"]).startswith(p) for p in pref)]
    unplaced = [r["acq_id"] for r in sel if d3_relpath(r) is None]
    if unplaced:
        raise WorkbookError(f"{len(unplaced)} drive-3 acquisition(s) with a blank project have no drive path in "
                            f"original_name or notes (e.g. {unplaced[0]}); refusing to guess their group.")
    rel = {r["acq_id"]: d3_relpath(r) for r in sel}
    cl = claims_for(set(rel.values()), claims_path) if sel else {}
    out = []
    for r in sel:
        verdict, claim_id, proposed = cl[rel[r["acq_id"]]]
        out.append({"acq_id": r["acq_id"],
                    "batch": os.path.splitext(os.path.basename(norm_config(r["ingest_config"])))[0],
                    "instrument": r["instrument"], "acquisition_date": r["acquisition_datetime"][:10],
                    "verdict": verdict, "path_claimed_project": proposed, "claim_id": claim_id,
                    "researcher": "", "drive": D3, "relpath": rel[r["acq_id"]], "archive": "", "member": "",
                    "other_copies": ""})
    return out


def common_folder(folders):
    parts = [f.split("\\") for f in folders]
    if not parts:
        return ""
    common = parts[0]
    for p in parts[1:]:
        n = 0
        while n < min(len(common), len(p)) and common[n] == p[n]:
            n += 1
        common = common[:n]
    return "\\".join(common)


def plan_assign(rows, wb, manifest_rows, acqs, today, source):
    """(appends, readme line, report) for assign-append."""
    import nonraw_placement as NP
    ws_g, ws_a = sheet(wb, "Groups to assign"), sheet(wb, "Acquisitions")
    hg = need(ws_g, ["Group", "Priority", "Researcher folder", "Drive", "Folder on the drive", "Raw acquisitions",
                     "Other files", "Other files (GB)", "From", "To", "Example file names",
                     "Where the other files are now", "Raw acquisitions (first .. last ACQ-ID)", "group key (do not edit)"])
    ha = need(ws_a, ["Group", "ACQ-ID", "Acquired", "Instrument", "File name", "Folder on the drive", "Drive",
                     "Researcher folder", "Claimed project (uncertain)"])
    g_by_key = {}
    max_g = 0
    for r in range(2, last_row(ws_g) + 1):
        g, key = ws_g.cell(row=r, column=hg["Group"]).value, ws_g.cell(row=r, column=hg["group key (do not edit)"]).value
        if key:
            g_by_key[NP.norm_key(key)] = g
        m = re.match(r"^G(\d+)$", str(g or ""))
        if m:
            max_g = max(max_g, int(m.group(1)))
    have_acqs = {v for v in column_values(ws_a, ha["ACQ-ID"]) if v}

    groups = collections.OrderedDict()

    def grp(gk):
        d, who, series, kind = gk
        key = NP.key_string(d, who, series)
        return groups.setdefault(key, {"key": key, "researcher": who, "series": series, "kind": kind,
                                       "files": [], "bytes": 0, "folders": [], "acqs": []})
    n_hold = 0
    for r in manifest_rows:
        if r["drive"] != D3 or r["decision"] != "holding" or not NP.mappable(r["reason"]):
            continue
        n_hold += 1
        g = grp(NP.manifest_group(r))
        g["files"].append(r["member"] or r["relpath"])
        g["bytes"] += int(r["size"] or 0)
        g["folders"].append(os.path.dirname(r["dest_rel"].replace("/", "\\")))
    for a in acqs:
        grp(NP.blank_list_group(a))["acqs"].append(a)

    new = [g for g in groups.values() if NP.norm_key(g["key"]) not in g_by_key]
    new.sort(key=lambda g: (not (g["bytes"] >= 1e9 or len(g["acqs"]) >= 20), -g["bytes"], -len(g["acqs"]), g["key"]))
    for i, g in enumerate(new, start=max_g + 1):
        g["G"] = f"G{i:03d}"
        g_by_key[NP.norm_key(g["key"])] = g["G"]
    g_rows = []
    for g in new:
        ids = sorted(a["acq_id"] for a in g["acqs"])
        dates = sorted(a["acquisition_date"] for a in g["acqs"] if a["acquisition_date"])
        ex = sorted({os.path.basename(f.replace("/", "\\")) for f in g["files"]})[:3]
        g_rows.append({
            "Group": g["G"], "Priority": "A" if g["bytes"] >= 1e9 or len(ids) >= 20 else "B",
            "Researcher folder": g["researcher"] or None, "Drive": D3_DISPLAY, "Folder on the drive": g["series"],
            "Raw acquisitions": len(ids), "Other files": len(g["files"]),
            "Other files (GB)": round(g["bytes"] / 1e9, 2),
            "From": dates[0] if dates else None, "To": dates[-1] if dates else None,
            "Example file names": "\n".join(ex) or None,
            "Where the other files are now": (DATA_SHARE + "\\" + common_folder(g["folders"])) if g["files"] else "(no other files)",
            "Raw acquisitions (first .. last ACQ-ID)": (";".join(ids) if len(ids) <= 25 else
                                                       f"{len(ids)} ACQ-IDs: {ids[0]} .. {ids[-1]}") or None,
            "group key (do not edit)": g["key"], ADDED: f"{today}, {source}"})

    new_acqs = []
    for a in acqs:
        if a["acq_id"] in have_acqs:
            continue
        d, who, series, _k = NP.blank_list_group(a)
        gid = g_by_key[NP.norm_key(NP.key_string(d, who, series))]
        folder, _sep, fname = a["relpath"].rpartition("\\")
        new_acqs.append((gid, a, folder, fname))
    new_acqs.sort(key=lambda t: (t[0], t[1]["acq_id"]))
    a_rows = [{"Group": gid, "ACQ-ID": a["acq_id"], "Acquired": a["acquisition_date"], "Instrument": a["instrument"],
               "File name": fname, "Folder on the drive": folder or None, "Drive": D3_DISPLAY,
               "Researcher folder": a["researcher"] or None,
               "Claimed project (uncertain)": (a["path_claimed_project"] or None) if a["verdict"] == "C" else None,
               ADDED: f"{today}, {source}"} for gid, a, folder, fname in new_acqs]

    in_listed = sum(1 for gid, _a, _f, _n in new_acqs if gid not in {g["G"] for g in new})
    line = None
    if g_rows or a_rows:
        span = f", groups {new[0]['G']}–{new[-1]['G']}" if new else ""
        line = (f"Added {today}: {len(new):,} groups from the {source} ({D3_DISPLAY}){span}: "
                f"{sum(len(g['files']) for g in new):,} other files "
                f"({sum(g['bytes'] for g in new) / 1e9:,.1f} GB) and {sum(len(g['acqs']) for g in new):,} raw "
                f"acquisitions without a project"
                + (f"; and {in_listed:,} more raw acquisitions of groups already listed" if in_listed else "")
                + f". They are marked in the column 'Added'. Their other files are in the holding folder "
                  f"{DATA_SHARE}\\staging\\historical_drives_unassigned\\{D3_DISPLAY}\\ .")
    report = {
        "assign_append_groups.csv": (
            ["group", "group_key", "kind", "priority", "raw_acqs", "other_files", "other_gb"],
            [{"group": r["Group"], "group_key": r["group key (do not edit)"], "kind": g["kind"],
              "priority": r["Priority"], "raw_acqs": r["Raw acquisitions"], "other_files": r["Other files"],
              "other_gb": r["Other files (GB)"]} for r, g in zip(g_rows, new)]),
        "drive3_blank_project_list.csv": (BLANK_LIST_FIELDS, [a for _g, a, _f, _n in new_acqs]),
    }
    print(f"drive-3 holding files that 2b can map: {n_hold:,}; acquisitions with a blank project: {len(acqs):,}")
    print(f"groups: {len(groups):,} in all, {len(new):,} new; acquisition rows to append: {len(a_rows):,} "
          f"({in_listed:,} of them in groups already listed)")
    return [SheetAppend("Groups to assign", g_rows), SheetAppend("Acquisitions", a_rows)], line, report


def cmd_assign(args):
    workbook = args.workbook or os.path.join(args.nas_root, "projects", ASSIGN_NAME)
    for p, what in ((args.manifest, "--manifest"), (args.claims, "--claims")):
        if not os.path.isfile(p):
            raise WorkbookError(f"{what} not found: {p}")
    rows, _projects = read_registry(args.nas_root)
    acqs = blank_project_acqs(rows, args.config_prefixes or D3_PREFIXES, args.claims)
    data, wb = load_target(workbook)
    appends, line, report = plan_assign(rows, wb, read_csv(args.manifest), acqs, args.date, args.source)
    return run_write("assign-append", workbook, data, wb, appends, line, report, args)


# ---- CLI ----------------------------------------------------------------------------

def main(argv=None):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass
    parser = argparse.ArgumentParser(description="Append to the claim / assign workbooks; never rebuild them. "
                                                 "Dry run unless --apply.")
    sub = parser.add_subparsers(dest="cmd", required=True)
    for name in ("claims-append", "assign-append"):
        p = sub.add_parser(name)
        p.add_argument("--nas-root", default=os.environ.get("GJESUS3_ROOT", r"J:\gjesus3-data"),
                       help="the NAS root (read: registries\\; the default workbook is under projects\\)")
        p.add_argument("--workbook", help="the workbook (default: the shared one under <nas-root>\\projects\\)")
        p.add_argument("--date", default=dt.date.today().isoformat(), help="the date in the Added column")
        p.add_argument("--apply", action="store_true", help="write the workbook (default: a preview copy only)")
        p.add_argument("--backup-dir", help="--apply: a folder that does not exist yet, off the NAS")
        p.add_argument("--preview-dir", default=os.path.join(tempfile.gettempdir(), "gjesus3_claim_workbooks"),
                       help="dry run: where the preview copy and the report go (never beside the workbook)")
        if name == "assign-append":
            p.add_argument("--manifest", default=D3_MANIFEST, help="the drive-3 placement manifest (gated v2)")
            p.add_argument("--claims", default=D3_CLAIMS, help="part A2's file_claims.csv for drive 3")
            p.add_argument("--config-prefix", dest="config_prefixes", action="append", default=[],
                           help=f"configs whose blank-project acquisitions are listed (repeatable; default "
                                f"{', '.join(D3_PREFIXES)})")
            p.add_argument("--source", default=D3_SOURCE, help="the source named in the Added column")
    args = parser.parse_args(argv)
    if args.apply and not args.backup_dir:
        log("--apply needs --backup-dir <a folder that does not exist yet, off the NAS>.", "ERROR")
        return 2
    if not re.match(r"^\d{4}-\d{2}-\d{2}$", args.date):
        log("--date must be YYYY-MM-DD.", "ERROR")
        return 2
    try:
        return cmd_claims(args) if args.cmd == "claims-append" else cmd_assign(args)
    except WorkbookError as exc:
        log(str(exc), "ERROR")
        return 1


if __name__ == "__main__":
    sys.exit(main())
